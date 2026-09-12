import csv
import datetime
import json
import re
import urllib.request

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

GITHUB_API_URL = "https://api.github.com/repos/andreanini/jacktherippercorpus/contents/corpus"
OUTPUT_CSV = "jack_the_ripper_corpus.csv"
OUTPUT_XLSX = "jack_the_ripper_corpus.xlsx"


def extract_date(filename: str, text: str) -> datetime.date | None:
    """
    Parses dates into datetime.date objects for sorting.
    """
    # 1. Match DDMMYY where year YY is 88-92 (1888-1892)
    m_ddmmyy = re.search(r'(?<!\d)(\d{2})(\d{2})(8[89]|9[0-2])(?!\d)', filename)
    if m_ddmmyy:
        d, m, y = m_ddmmyy.groups()
        day, month, year = int(d), int(m), 1800 + int(y)
        if 1 <= month <= 12 and 1 <= day <= 31:
            return datetime.date(year, month, day)

    # 2. Match YYYY-MM-DD or YYYY_MM_DD
    m_iso = re.search(r'(18\d{2})[-_](\d{2})[-_](\d{2})', filename)
    if m_iso:
        y, m, d = m_iso.groups()
        year, month, day = int(y), int(m), int(d)
        if 1 <= month <= 12 and 1 <= day <= 31:
            return datetime.date(year, month, day)

    # 3. Match YYYYMMDD
    m_compact = re.search(r'(18\d{2})(\d{2})(\d{2})', filename)
    if m_compact:
        y, m, d = m_compact.groups()
        year, month, day = int(y), int(m), int(d)
        if 1 <= month <= 12 and 1 <= day <= 31:
            return datetime.date(year, month, day)

    # 4. Fallback: Search text header for DD/MM/YYYY or DD-MM-YYYY
    m_text = re.search(r'\b(\d{1,2})[/-](\d{1,2})[/-](18\d{2})\b', text[:300])
    if m_text:
        d, m, y = m_text.groups()
        day, month, year = int(d), int(m), int(y)
        if 1 <= month <= 12 and 1 <= day <= 31:
            return datetime.date(year, month, day)

    return None


def download_sort_and_process():
    headers = {"User-Agent": "Mozilla/5.0"}
    req = urllib.request.Request(GITHUB_API_URL, headers=headers)

    print("Fetching file list from GitHub repository...")
    with urllib.request.urlopen(req) as response:
        file_list = json.loads(response.read().decode("utf-8"))

    corpus_data = []

    for file_info in file_list:
        if file_info.get("type") == "file" and file_info["name"].endswith(".txt"):
            filename = file_info["name"]
            download_url = file_info["download_url"]

            file_req = urllib.request.Request(download_url, headers=headers)
            with urllib.request.urlopen(file_req) as file_resp:
                raw_text = file_resp.read().decode("utf-8", errors="ignore")

            date_obj = extract_date(filename, raw_text)
            date_str = date_obj.strftime("%d/%m/%Y") if date_obj else "UNKNOWN"

            corpus_data.append({
                "filename": filename,
                "date_obj": date_obj,
                "date_str": date_str,
                "text": raw_text.strip()
            })

    # Sort chronologically by date; place UNKNOWN dates at the end
    corpus_data.sort(key=lambda x: (x["date_obj"] is None, x["date_obj"]))

    print(f"\nSuccessfully sorted {len(corpus_data)} files chronologically.")

    # 1. Export CSV
    with open(OUTPUT_CSV, mode="w", newline="", encoding="utf-8-sig") as csv_file:
        writer = csv.writer(csv_file)
        writer.writerow(["filename", "date", "text"])
        for row in corpus_data:
            formatted_date = f'="{row["date_str"]}"' if row["date_str"] != "UNKNOWN" else "UNKNOWN"
            writer.writerow([row["filename"], formatted_date, row["text"]])
    print(f"Saved sorted CSV: '{OUTPUT_CSV}'")

    # 2. Export XLSX (Text-formatted cells to avoid pre-1900 Excel date bug)
    if HAS_OPENPYXL:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Corpus"

        ws.append(["filename", "date", "text"])

        for row_idx, row in enumerate(corpus_data, start=2):
            ws.cell(row=row_idx, column=1, value=row["filename"])
            
            date_cell = ws.cell(row=row_idx, column=2, value=row["date_str"])
            date_cell.number_format = '@'  # Explicit text format

            ws.cell(row=row_idx, column=3, value=row["text"])

        wb.save(OUTPUT_XLSX)
        print(f"Saved sorted Excel workbook: '{OUTPUT_XLSX}'")


if __name__ == "__main__":
    download_sort_and_process()